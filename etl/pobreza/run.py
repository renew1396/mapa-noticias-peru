"""ETL de pobreza monetaria distrital (INEI, Mapa de Pobreza Monetaria Provincial y Distrital 2018).

Es la edición más reciente que existe a nivel distrital (publicada en febrero de 2020). El INEI publica un
INTERVALO de confianza al 95 % (inferior–superior), no un valor único: la web muestra el intervalo y colorea
por el punto medio, dejando claro que es una estimación.

Salidas:
  data/layers/pobreza.json        {ubigeo: [inferior, superior, punto_medio, grupo_robusto, puesto, poblacion_2020]}
  data/layers/pobreza.meta.json   incluye los cortes de las 5 clases (quintiles del punto medio)

Uso:  python -m etl.pobreza.run
"""
from __future__ import annotations

import io
import json
import sys

import openpyxl

from etl.common import EtlError, LAYERS_DIR, RAW_DIR, ROOT, http_get, ubigeo6, write_json, write_meta

URL = "https://cdn.www.gob.pe/uploads/document/file/3340937/Anexo%20Estad%C3%ADstico.xlsx"
PAGINA = "https://www.gob.pe/institucion/inei/informes-publicaciones/3204872-mapa-de-pobreza-provincial-y-distrital-2018"
DISTRITOS_ESPERADOS = 1874


def main() -> int:
    contenido = http_get(URL, timeout=300).content
    (RAW_DIR / "pobreza").mkdir(parents=True, exist_ok=True)
    (RAW_DIR / "pobreza" / "anexo_estadistico_2018.xlsx").write_bytes(contenido)
    wb = openpyxl.load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    for hoja in ("Anexo1", "Anexo2"):
        if hoja not in wb.sheetnames:
            raise EtlError(f"El Excel del INEI ya no tiene la hoja {hoja!r}. Hojas: {wb.sheetnames}")

    # Anexo 2 trae los 1 874 distritos (filas con "sufijo" numérico); el Anexo 1 omite uno (150906), así que
    # de él solo se toma el grupo robusto.
    grupos = {}
    for r in wb["Anexo1"].iter_rows(values_only=True):
        if r and r[0] is not None and str(r[0]).strip().isdigit() and len(r) > 7:
            grupos[ubigeo6(r[0])] = int(r[7]) if str(r[7]).strip().isdigit() else None

    datos: dict[str, list] = {}
    for r in wb["Anexo2"].iter_rows(values_only=True):
        if not r or r[0] is None or not str(r[0]).strip().isdigit() or not str(r[1] or "").strip().isdigit():
            continue
        ub = ubigeo6(r[0])
        pob, inf, sup, puesto = r[3], r[4], r[5], r[6]
        if not all(isinstance(x, (int, float)) for x in (inf, sup)) or not 0 <= inf <= sup <= 100:
            raise EtlError(f"Fila con intervalo inválido para {ub}: {r}")
        datos[ub] = [round(inf, 1), round(sup, 1), round((inf + sup) / 2, 1), grupos.get(ub),
                     int(puesto) if isinstance(puesto, (int, float)) else None,
                     int(pob) if isinstance(pob, (int, float)) else None]
    if len(datos) != DISTRITOS_ESPERADOS:
        raise EtlError(f"Se esperaban {DISTRITOS_ESPERADOS} distritos y llegaron {len(datos)}: ¿cambió el formato?")

    # Cortes de 5 clases por quintiles del punto medio (los usa la leyenda de la web).
    medios = sorted(v[2] for v in datos.values())
    cortes = [medios[int(len(medios) * q / 5)] for q in range(1, 5)]

    actuales = {d["ubigeo"] for d in json.loads((ROOT / "data/gazetteer/ubigeos.json").read_text(encoding="utf-8"))}
    sin_dato = sorted(actuales - set(datos))
    write_json(LAYERS_DIR / "pobreza.json", dict(sorted(datos.items())))
    write_meta(
        "pobreza",
        fuente="INEI – Mapa de Pobreza Monetaria Provincial y Distrital 2018",
        url=PAGINA,
        fecha_corte="2018",
        visualizacion="coropletico",
        registros=len(datos),
        notas=("Intervalo de confianza al 95 % de la pobreza monetaria total (% de la población). El color usa el "
               f"punto medio del intervalo. {len(sin_dato)} distritos actuales no tienen dato porque se crearon "
               f"después del mapa: {', '.join(sin_dato)}. Campos: [inferior, superior, punto_medio, grupo_robusto, "
               "puesto_nacional (1 = más pobre), poblacion_2020]."),
        estilo={"cortes": cortes, "colores": ["#fff7bc", "#fec44f", "#fe9929", "#d95f0e", "#993404"],
                "sin_dato": "#e5e7eb"},
    )
    print(f"OK: {len(datos)} distritos con pobreza 2018; cortes {cortes}; {len(sin_dato)} distritos actuales sin dato")
    return 0


if __name__ == "__main__":
    sys.exit(main())
